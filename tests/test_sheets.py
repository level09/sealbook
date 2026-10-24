"""Spreadsheet intake: parsing, cell spans, dates, mapping memory and suggestions,
version diffs, and the extract_sheet run against a throwaway Postgres. The model and
the embedding server are mocked at their HTTP and function boundaries."""

import io
import os
import unittest
from datetime import UTC, date, datetime
from pathlib import Path
from unittest import mock

import httpx
import orjson
from openpyxl import Workbook
from sqlalchemy import select

import stk.extensions as ext
from stk import blobs
from stk.ledger.contracts import Prop, Tier
from stk.ledger.models import Acquisition, Blob, Run, Statement, StatementTier
from stk.ledger.span import validate
from stk.ledger.write import record_acquisition, verify_db
from stk.runs.worker import work_once
from stk.settings import Config
from stk.sheets import dates, mapping, parse
from stk.sheets.extract import ROWS_PER_CALL, pin, value_of
from stk.sheets.models import Dataset, DatasetVersion
from tests import test_review

DATA = Path(__file__).resolve().parent.parent / "acceptance" / "data"
V1 = (DATA / "messy_people.csv").read_bytes()
V2 = (DATA / "messy_people_v2.csv").read_bytes()
HEADERS = parse.read(V1, parse.CSV)[0].rows[0]
MAPPING = dict(
    zip(HEADERS, ["key", "person", "event", "place", "date", "ignore"], strict=True)
)


def xlsx() -> bytes:
    book = Workbook()
    ws = book.active
    ws.title = "People"
    ws.append(["id", "name", "seen"])
    ws.append([1, "Rania Haddad", datetime(2014, 3, 15)])
    ws.append([2.0, None, 41.5])
    book.create_sheet("Empty")
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


class ParseTest(unittest.TestCase):
    def test_csv_bom_semicolons_and_padding(self):
        data = "﻿a;b;c\n1;x\n\n".encode()
        (sheet,) = parse.read(data, parse.CSV)
        self.assertEqual(sheet.rows, (("a", "b", "c"), ("1", "x", "")))
        with self.assertRaisesRegex(ValueError, "UTF-8"):
            parse.read("a,b\nع,1".encode("cp1256"), parse.CSV)

    def test_xlsx_values_as_text(self):
        people, empty = parse.read(xlsx(), parse.XLSX)
        self.assertEqual(people.name, "People")
        self.assertEqual(
            people.rows,
            (
                ("id", "name", "seen"),
                ("1", "Rania Haddad", "2014-03-15"),
                ("2", "", "41.5"),
            ),
        )
        self.assertEqual(empty.rows, ())
        with self.assertRaisesRegex(ValueError, "XLSX"):
            parse.read(b"not a zip", parse.XLSX)

    def test_mime_from_extension(self):
        self.assertEqual(
            parse.mime_for("List.CSV", "application/vnd.ms-excel"), parse.CSV
        )
        self.assertEqual(parse.mime_for("a.xlsx", ""), parse.XLSX)
        self.assertEqual(parse.mime_for("a.txt", "text/plain"), "text/plain")
        self.assertEqual([parse.letter(n) for n in (1, 26, 27)], ["A", "Z", "AA"])


class CellSpanTest(unittest.TestCase):
    def test_round_trip_and_bounds(self):
        sheets = parse.read(V1, parse.CSV)
        meta = {"sheets": parse.lengths(V1, parse.CSV)}
        span = pin(sheets[0], 1, 2, 2, "سامر خليل النجار")
        self.assertEqual(
            span, {"sheet": 1, "row": 2, "column": 2, "start": 0, "end": 16}
        )
        self.assertEqual(validate(span, meta), (True, ""))
        self.assertEqual(parse.cell(sheets, span), "سامر خليل النجار")
        whole = pin(sheets[0], 1, 2, 4, "حرستا")
        self.assertEqual(whole, {"sheet": 1, "row": 2, "column": 4})
        self.assertEqual(parse.cell(sheets, whole), "حرستا")
        # The model named the wrong column: the quote is found in its own cell.
        self.assertEqual(pin(sheets[0], 1, 2, 4, "نجار")["column"], 2)
        self.assertIsNone(pin(sheets[0], 1, 2, 2, "not in the row"))
        cases = [
            ({"sheet": 2, "row": 1, "column": 1}, "sheet_out_of_range"),
            ({"sheet": 1, "row": 17, "column": 1}, "cell_out_of_range"),
            ({"sheet": 1, "row": 1, "column": 7}, "cell_out_of_range"),
            ({"sheet": 1, "row": True, "column": 1}, "invalid_cell"),
            (
                {"sheet": 1, "row": 2, "column": 2, "start": 0, "end": 99},
                "char_range_out_of_bounds",
            ),
            ({"sheet": 1, "row": 2, "column": 2, "start": 0}, "invalid_span_keys"),
        ]
        for span, why in cases:
            self.assertEqual(validate(span, meta), (False, why), span)
        self.assertEqual(
            validate({"sheet": 1, "row": 1, "column": 1}, {"sheets": None}),
            (False, "sheets_required"),
        )


class DatesTest(unittest.TestCase):
    def test_table(self):
        d = date
        table = {
            "2013-05-12": (d(2013, 5, 12), d(2013, 5, 12)),
            "15/3/2013": (d(2013, 3, 15), d(2013, 3, 15)),
            "3/15/2013": (d(2013, 3, 15), d(2013, 3, 15)),
            "١٢/١١/٢٠١٣": (d(2013, 11, 12), d(2013, 11, 12)),
            "٢٠١٤/٠٧/٠٣": (d(2014, 7, 3), d(2014, 7, 3)),
            "اذار 2011": (d(2011, 3, 1), d(2011, 3, 31)),
            "آذار ٢٠١١": (d(2011, 3, 1), d(2011, 3, 31)),
            "شباط 2012": (d(2012, 2, 1), d(2012, 2, 29)),
            "تشرين الأول 2012": (d(2012, 10, 1), d(2012, 10, 31)),
            "July 2013": (d(2013, 7, 1), d(2013, 7, 31)),
            "15 March 2013": (d(2013, 3, 15), d(2013, 3, 15)),
            "5/2013": (d(2013, 5, 1), d(2013, 5, 31)),
            "2014": (d(2014, 1, 1), d(2014, 12, 31)),
            "around 2016": (d(2016, 1, 1), d(2016, 12, 31)),
            "منذ 2012": (d(2012, 1, 1), d(2012, 12, 31)),
            "2012-2013": (d(2012, 1, 1), d(2013, 12, 31)),
            "unknown": None,
            "غير معروف": None,
            "": None,
            "31/2/2013": None,
            "soon": None,
        }
        for text, want in table.items():
            self.assertEqual(dates.interval(text), want, text)
        self.assertTrue(dates.unknown("Unknown"))
        self.assertFalse(dates.unknown("soon"))


class MappingTest(unittest.IsolatedAsyncioTestCase):
    async def test_remembered_columns_skip_the_model(self):
        cols = mapping.columns(parse.read(V1, parse.CSV)[0])
        self.assertEqual(
            cols[1]["samples"][0], "سامر خليل النجار - ذكر - 29 سنة - نجار"
        )
        self.assertEqual(len(cols[1]["samples"]), 5)
        embed = mock.AsyncMock()
        with mock.patch.object(mapping, "embed", embed):
            out = await mapping.suggest(
                cols, {h.upper(): f for h, f in MAPPING.items()}
            )
        embed.assert_not_awaited()
        self.assertEqual([c["remembered"] for c in out], list(MAPPING.values()))

    async def test_unfamiliar_columns_are_ranked(self):
        cols = mapping.columns(parse.read(V1, parse.CSV)[0])
        n = len(mapping.FIELDS)

        async def embed(texts):
            # Fields are one-hot; each column points at the field of its position.
            self.assertEqual(len(texts), n + 2)
            self.assertIn("سامر خليل النجار", texts[n])
            vectors = [[float(i == j) for j in range(n)] for i in range(n)]
            keys = [f.key for f in mapping.FIELDS]
            for field in ("person", "date"):
                v = [0.0] * n
                v[keys.index(field)] = 0.8
                v[keys.index("name")] = 0.6
                vectors.append(v)
            return vectors

        remembered = {h: f for h, f in MAPPING.items() if f not in ("person", "date")}
        with mock.patch.object(mapping, "embed", embed):
            out = await mapping.suggest(cols, remembered)
        person, dated = out[1], out[4]
        self.assertIsNone(person["remembered"])
        self.assertEqual(
            [s["field"] for s in person["suggestions"]][:2], ["person", "name"]
        )
        self.assertEqual(person["suggestions"][0]["score"], 0.8)
        self.assertEqual(dated["suggestions"][0]["field"], "date")
        self.assertEqual(out[0]["suggestions"], [])

    def test_check(self):
        sheet = parse.read(V1, parse.CSV)[0]
        mapping.check(sheet, "No.", MAPPING)
        for key, cols, why in (
            ("No.", {**MAPPING, "No.": "person"}, "key column"),
            (
                "No.",
                {k: v for k, v in MAPPING.items() if k != "التاريخ"},
                "every column",
            ),
            ("No.", {**MAPPING, "التاريخ": "birthday"}, "unknown fields"),
        ):
            with self.assertRaisesRegex(ValueError, why):
                mapping.check(sheet, key, cols)


class DiffTest(unittest.TestCase):
    def test_versions(self):
        v1, v2 = (parse.read(b, parse.CSV)[0] for b in (V1, V2))
        first = mapping.diff(v1, None, "No.", MAPPING)
        self.assertEqual(mapping.counts(first)["new"], 15)
        rows = mapping.diff(v2, v1, "No.", MAPPING)
        self.assertEqual(
            mapping.counts(rows),
            {"new": 2, "changed": 3, "unchanged": 10, "missing": 1, "invalid": 3},
        )
        by = {(r["status"], r["key"]): r for r in rows}
        self.assertEqual(by["changed", "2"]["changed"], ["التاريخ"])
        self.assertEqual(by["changed", "15"]["changed"], ["المكان"])
        # Only the ignored source note changed on row 8.
        self.assertIn(("unchanged", "8"), by)
        self.assertEqual(by["missing", "9"]["previous_row"], 10)
        invalid = sorted((r["row"], r["why"]) for r in rows if r["status"] == "invalid")
        self.assertEqual(
            invalid, [(12, "duplicate_key"), (13, "duplicate_key"), (19, "blank_key")]
        )
        # A key present only on invalid rows is not reported missing.
        self.assertNotIn(("missing", "12"), by)


class ValueTest(unittest.TestCase):
    def test_normalised_values(self):
        fact = {"quote": "ذكر", "lang": "ar", "norm": "Male", "kind": None}
        self.assertEqual(value_of(Prop.SEX, fact)["norm"], "male")
        self.assertIsNone(value_of(Prop.SEX, fact | {"norm": "unknown"}))
        self.assertEqual(value_of(Prop.AGE, fact | {"norm": "٣٨"})["norm"], 38)
        self.assertNotIn("norm", value_of(Prop.AGE, fact | {"norm": "about"}))
        self.assertIsNone(value_of(Prop.EVENT, fact | {"kind": "arrested"}))
        self.assertEqual(
            value_of(Prop.EVENT, fact | {"kind": "killed"})["kind"], "killed"
        )


def completion(rows):
    content = orjson.dumps({"rows": rows}).decode()
    return {
        "model": "gemma-local",
        "choices": [{"finish_reason": "stop", "message": {"content": content}}],
        "usage": {"prompt_tokens": 1, "completion_tokens": 1},
    }


def fact(prop, column, quote, norm=None, kind=None, when=None, lang="ar"):
    return {
        "prop": prop,
        "column": column,
        "quote": quote,
        "lang": lang,
        "norm": norm,
        "kind": kind,
        "date": when,
    }


ROW2 = [
    fact("name", 2, "سامر خليل النجار"),
    fact("sex", 2, "ذكر", norm="male"),
    fact("age", 2, "29 سنة", norm="29"),
    fact("occupation", 2, "نجار"),
    fact("arrested_at", 4, "حرستا", when="2013-05-12"),
    fact("transferred_to", 3, "الفرع 999", when="2013-05-12"),
    fact("seen_at", 4, "Paris", when="2013-05-12"),
]
ROW4 = [
    fact("event", 3, "قتل بقصف على السوق", kind="killed", when="اذار 2013"),
    fact("event", 3, "قتل", kind="killed", when="1999"),
]


class SheetFlowTest(unittest.IsolatedAsyncioTestCase):
    setUp_review = test_review.ReviewTest.setUp
    asyncSetUp = test_review.ReviewTest.asyncSetUp
    asyncTearDown = test_review.ReviewTest.asyncTearDown
    client = test_review.ReviewTest.client

    def setUp(self):
        self.setUp_review()
        self.enterContext(
            mock.patch.object(Config, "RUN_MODELS", {"local": "gemma-local"})
        )
        self.enterContext(
            mock.patch.dict(
                os.environ, {"ROUTE_LOCAL_BASE_URL": "http://local.test/v1"}
            )
        )
        self.sent = []

    def model(self, rows):
        def handle(request):
            self.sent.append(orjson.loads(request.content))
            return httpx.Response(200, json=completion(rows))

        return httpx.MockTransport(handle)

    async def sheet(self, data, sensitivity="ordinary", name="messy_people.csv"):
        async with ext.async_session_factory() as s, s.begin():
            blob = await blobs.put(s, data, parse.CSV, Config.BLOB_DIR)
            acq = await record_acquisition(
                s,
                blob_sha256=blob.sha256,
                origin="device:test",
                collector_id=self.agent.id,
                at=datetime.now(UTC),
                context={"investigation": "case-sheet", "filename": name},
                sensitivity=sensitivity,
                doc_type="register",
            )
        return blob.sha256, acq.id

    async def test_confirm_extract_and_second_version(self):
        sha, acq = await self.sheet(V1)
        client = await self.client()
        with mock.patch.object(mapping, "embed", side_effect=ValueError("down")):
            r = await client.get(f"/api/sheets/{acq}/preview")
        body = await r.get_json()
        self.assertEqual(r.status_code, 200, body)
        self.assertEqual(body["table"]["total"], 15)
        self.assertEqual(body["suggest_error"], "down")
        item = {
            "dataset_name": "Missing people list",
            "key_column": "No.",
            "mapping": MAPPING,
        }
        r = await client.post(
            f"/api/sheets/{acq}/diff", json={"item": item}, headers=client.csrf
        )
        self.assertEqual((await r.get_json())["counts"]["new"], 15)
        async with ext.async_session_factory() as s:
            self.assertIsNone(await s.scalar(select(Dataset.id)))  # diff writes nothing
        r = await client.post(
            f"/api/sheets/{acq}/confirm", json={"item": item}, headers=client.csrf
        )
        out = await r.get_json()
        self.assertEqual((r.status_code, out["queued"]), (200, "extract_sheet"), out)
        again = await client.post(
            f"/api/sheets/{acq}/confirm", json={"item": item}, headers=client.csrf
        )
        self.assertEqual(again.status_code, 409)

        transport = self.model([{"row": 2, "facts": ROW2}, {"row": 4, "facts": ROW4}])
        self.assertTrue(
            await work_once(ext.async_session_factory, ["extract_sheet"], transport)
        )
        sent = self.sent[0]
        self.assertEqual(sent["response_format"]["type"], "json_schema")
        self.assertEqual(sent["reasoning_effort"], "none")  # local route
        cells = orjson.loads(sent["messages"][1]["content"])
        self.assertEqual(len(cells), ROWS_PER_CALL)  # first chunk of rows
        self.assertNotIn(
            "اتصال هاتفي", sent["messages"][1]["content"]
        )  # ignored column
        async with ext.async_session_factory() as s:
            run = await s.get(Run, out["run"])
            self.assertEqual(run.status, "done", run.error)
            self.assertEqual(run.inputs["mapping"], MAPPING)
            self.assertEqual(run.model_served, "gemma-local")
            raw = orjson.loads(
                blobs.read(await s.get(Blob, run.raw_output_ref), Config.BLOB_DIR)
            )
            self.assertEqual(
                len(raw["raw"]), -(-15 // ROWS_PER_CALL)
            )  # one model output per chunk of rows
            rows = (
                await s.execute(
                    select(Statement, StatementTier.tier, StatementTier.reason)
                    .join(StatementTier, StatementTier.statement_id == Statement.id)
                    .where(Statement.run_id == run.id)
                )
            ).all()
            by = {
                (st.prop, st.span and st.span.get("row")): (st, t, why)
                for st, t, why in rows
            }
            name, tier, _ = by["name", 2]
            self.assertEqual(
                name.span, {"sheet": 1, "row": 2, "column": 2, "start": 0, "end": 16}
            )
            self.assertEqual(name.span_blob, sha)
            self.assertEqual(
                tier, Tier.UNCONFIRMED
            )  # nothing is accepted automatically
            self.assertEqual(
                len({by[p, 2][0].subject for p in ("name", "sex", "age")}), 1
            )
            arrested = by["arrested_at", 2][0]
            self.assertEqual(
                (arrested.valid_from, arrested.valid_to), (date(2013, 5, 12),) * 2
            )
            killed = [st for st, *_ in rows if st.prop == "event"]
            self.assertEqual(
                sorted(((k.valid_from, k.valid_to) for k in killed), key=str),
                [(None, None), (date(2013, 3, 1), date(2013, 3, 31))],
            )  # a date the row does not contain is ignored
            _, tier, why = by["seen_at", None]
            self.assertEqual((tier, why), (Tier.REJECTED, "no_span"))
            self.assertEqual(await verify_db(s), [])
            match = await s.scalar(select(Run).where(Run.tool == "match"))
            self.assertEqual(len(match.inputs["statements"]), 1)

        r = await client.get(f"/api/sheets/{acq}/proposals")
        grouped = await r.get_json()
        self.assertEqual([g["row"] for g in grouped["rows"]], [2, 4])
        place = next(
            c for c in grouped["rows"][0]["claims"] if c["prop"] == "arrested_at"
        )
        self.assertIsNone(place["place"])  # no gazetteer in this database: stays open

        sha2, acq2 = await self.sheet(V2, name="messy_people_v2.csv")
        embed = mock.AsyncMock()
        with mock.patch.object(mapping, "embed", embed):
            r = await client.get(
                f"/api/sheets/{acq2}/preview?dataset={out['version']['dataset']}"
            )
        cols = (await r.get_json())["columns"]
        embed.assert_not_awaited()
        self.assertEqual([c["remembered"] for c in cols], list(MAPPING.values()))
        item2 = {
            "dataset": out["version"]["dataset"],
            "key_column": "No.",
            "mapping": MAPPING,
        }
        r = await client.post(
            f"/api/sheets/{acq2}/confirm", json={"item": item2}, headers=client.csrf
        )
        out2 = await r.get_json()
        self.assertEqual(out2["plan"]["counts"]["changed"], 3)
        self.assertEqual(out2["version"]["number"], 2)
        async with ext.async_session_factory() as s:
            run2 = await s.get(Run, out2["run"])
            self.assertEqual(run2.inputs["rows"], [3, 7, 16, 17, 18])
            versions = (await s.execute(select(DatasetVersion.number))).scalars().all()
            self.assertEqual(sorted(versions), [1, 2])

    async def test_access_rules(self):
        sha, acq = await self.sheet(V1, sensitivity="sensitive")
        client = await self.client()
        r = await client.get(f"/api/sheets/{acq}/preview")
        self.assertEqual(
            (r.status_code, (await r.get_json())["reason"]), (403, "reveal_required")
        )
        self.assertEqual(
            (await client.get(f"/api/evidence/{sha}/sheet")).status_code, 403
        )
        r = await client.post(
            f"/api/evidence/{sha}/reveal", headers=client.csrf, json={"item": {}}
        )
        self.assertEqual(r.status_code, 200)
        with mock.patch.object(mapping, "embed", side_effect=ValueError("down")):
            self.assertEqual(
                (await client.get(f"/api/sheets/{acq}/preview")).status_code, 200
            )
        table = await (await client.get(f"/api/evidence/{sha}/sheet")).get_json()
        self.assertEqual(table["header"][0], "No.")

        from quart_security import hash_password

        from stk.user.models import User

        async with ext.async_session_factory() as s, s.begin():
            s.add(
                User(
                    email="plain@example.com",
                    name="plain",
                    password=hash_password("TestPassword123!", app=self.app),
                    active=True,
                    confirmed_at=datetime.now(),
                )
            )
        outsider = await self.client("plain@example.com")
        for path in (f"/api/sheets/{acq}/preview", f"/api/evidence/{sha}/sheet"):
            self.assertEqual((await outsider.get(path)).status_code, 403, path)
        r = await outsider.post(
            f"/api/sheets/{acq}/confirm", json={"item": {}}, headers=outsider.csrf
        )
        self.assertEqual(r.status_code, 403)

    async def test_ordinary_sheet_logs_a_view_and_is_not_read_automatically(self):
        from stk.ingest.service import queue_processing

        sha, acq = await self.sheet(V1)
        async with ext.async_session_factory() as s, s.begin():
            out = await queue_processing(s, await s.get(Acquisition, acq), parse.CSV)
        self.assertIsNone(out["queued"])
        client = await self.client()
        self.assertEqual(
            (await client.get(f"/api/evidence/{sha}/sheet")).status_code, 200
        )
        async with ext.async_session_factory() as s:
            views = await s.scalar(
                select(Statement.id).where(
                    Statement.subject == f"blob:{sha}", Statement.prop == Prop.ACCESS
                )
            )
        self.assertIsNotNone(views)


if __name__ == "__main__":
    unittest.main()
